"""Membangun data.json dari file KMZ dan Excel.
Pemakaian: python tools/bangun_data.py FOLDER_SUMBER [FOLDER_KELUARAN]"""
import re, json, math, glob, os, sys, zipfile, difflib
import xml.etree.ElementTree as ET
from openpyxl import load_workbook
SUMBER=sys.argv[1] if len(sys.argv)>1 else 'data_sumber'
OUT=sys.argv[2] if len(sys.argv)>2 else '.'
D=SUMBER.rstrip('/')+'/'
def norm(s):
    s=s.lower(); s=re.sub(r'\b(di|sbr|sumber|ters|desa)\b\.?',' ',s); s=re.sub(r'\(.*?\)',' ',s)
    s=re.sub(r'\b(19|20)\d\d\b',' ',s); s=re.sub(r'\b(ok|sip|fix|fixs|top)\b',' ',s)
    return re.sub(r'[^a-z0-9]','',s)
def hav(a,b):
    R=6371000; p1,p2=math.radians(a[0]),math.radians(b[0]); dp=p2-p1; dl=math.radians(b[1]-a[1])
    h=math.sin(dp/2)**2+math.cos(p1)*math.cos(p2)*math.sin(dl/2)**2; return 2*R*math.asin(math.sqrt(h))
def area_ha(c):
    la0=sum(p[0] for p in c)/len(c); kx=111320*math.cos(math.radians(la0)); ky=110540
    xy=[(p[1]*kx,p[0]*ky) for p in c]; s=0
    for i in range(len(xy)): x1,y1=xy[i]; x2,y2=xy[(i+1)%len(xy)]; s+=x1*y2-x2*y1
    return abs(s)/2/10000
def coords(t):
    out=[]
    for tok in t.split():
        p=tok.split(',')
        if len(p)>=2:
            try: out.append([round(float(p[1]),5),round(float(p[0]),5)])
            except: pass
    return out
def kml(path):
    z=zipfile.ZipFile(path); root=ET.fromstring(z.read([n for n in z.namelist() if n.endswith('.kml')][0]))
    B,S,A=[],[],[]
    for pm in root.iter():
        if not pm.tag.endswith('Placemark'): continue
        nm=''
        for ch in pm:
            if ch.tag.endswith('name'): nm=(ch.text or '').strip()
        for g in pm.iter():
            tg=g.tag.split('}')[-1]
            if tg=='Point':
                c=coords(g.find('.//{*}coordinates').text); 
                if c: B.append([nm,c[0][0],c[0][1]])
            elif tg=='LineString':
                c=coords(g.find('.//{*}coordinates').text)
                if len(c)>1: S.append([nm,round(sum(hav(c[i],c[i+1]) for i in range(len(c)-1))),c])
            elif tg=='Polygon':
                c=coords(g.find('.//{*}outerBoundaryIs//{*}coordinates').text)
                if len(c)>2: A.append([nm,round(area_ha(c),2),c])
    return B,S,A
# --- pilih kmz terbaru per DI
best={}
for f in glob.glob(D+"*.kmz"):
    b=os.path.basename(f)
    if b.startswith(('BATAS','SUNGAI')): continue
    m=re.match(r'^\d+[A-Za-z]?\s+(.*?)\s+((?:19|20)\d\d)\s+\w+\.kmz$',b)
    nm,yr=(m.group(1),int(m.group(2))) if m else (b[:-4],2023)
    k=norm(nm)
    if k not in best or yr>best[k][1]: best[k]=(f,yr,nm)
print("DI dari kmz:",len(best))
# --- xlsx
XLSX=sorted([f for f in glob.glob(D+'*.xlsx') if not os.path.basename(f).startswith('~$')],key=os.path.getmtime)[-1]
print('Excel:',os.path.basename(XLSX))
wb=load_workbook(XLSX,read_only=True,data_only=True)
def num(x):
    try: return float(x)
    except: return 0.0
r5a=list(wb['5a - O'].iter_rows(values_only=True))
dij={}; cur=None; order=[]
for i,r in enumerate(r5a):
    if r[1]=='JURU PENGAIRAN' and i+2<len(r5a): cur=str(r5a[i+2][1]).strip().title()
    elif r[1] and str(r[1]).strip().upper().startswith('DI') and cur:
        n=re.sub(r'^DI\.?\s*','',str(r[1]).strip()); 
        if n not in dij: dij[n]=cur; order.append(n)
r5=list(wb['5 - O auto '].iter_rows(values_only=True))
def rowvals(r):
    v=[num(x) for x in (list(r)+[None]*30)[:30]]
    return dict(baku=v[2],padi=sum(v[i] for i in (5,8,11,14,17)),persiapan=sum(v[i] for i in (3,4,6,7,9,10,12,13,15,16)),
        tebu=sum(v[18:21]),palawija=sum(v[21:24]),total=v[24],bero=v[25],panen=v[26])
def bkey(n): return norm(n)+('~s' if re.search(r'\b(sbr|sumber)\b',n,re.I) else '')
def blocks(rows):
    tan={}; blk=None
    for r in rows+[[None]*30]:
        c=r[1]; s=str(c).strip() if c else ''
        if s.upper().startswith('DI'):
            blk=bkey(re.sub(r'^DI\.?\s*','',s)); tan[blk]={'own':rowvals(r) if num(r[2])>0 else None,'ch':[]}
        elif blk and s and num(r[2])>0 and not s.upper().startswith(('JURU','UPTD','JUMLAH','TOTAL')):
            v=rowvals(r); v['_n']=s.lower(); v['_o']=s; tan[blk]['ch'].append(v)
    return tan
T5=blocks(r5); T5a=blocks(r5a)
ALIAS={'karsonopds':['karsonopedes'],'suko~s':['sumbersuko']}
def xkeys(n): k=bkey(n); return [k]+ALIAS.get(k,[])
def tanam(n):
    for src in (T5,T5a):
        for k in xkeys(n):
            t=src.get(k)
            if not t: continue
            if t['own']: return t['own']
            ch=t['ch']
            if not ch: continue
            ter=[c for c in ch if c['_n'].startswith('ters')]
            if ter: ch=ter
            return {k2:sum(c[k2] for c in ch) for k2 in ch[0] if k2 not in('_n','_o')}
    return None
r6=list(wb['06 - O'].iter_rows(values_only=True))
deb={}; curdi=None; nmr=None
for i,r in enumerate(r6):
    if r[1] and str(r[1]).isupper() and str(r[1]) not in('JURU PENGAIRAN','UPTD','PURWOSARI'): curdi=str(r[1])
    if r[2] and str(r[2]).strip(): nmr=str(r[2]).strip(); typ,leb,kon=r[4],r[5],r[6]
    if r[8]=='Q->' and curdi and nmr and i>0 and r6[i-1][8]=='H->':
        q=[num(x) for x in r[9:19] if isinstance(x,(int,float))]; h=[num(x) for x in r6[i-1][9:19] if isinstance(x,(int,float))]
        if q and nmr not in('Limpasan',): deb.setdefault(curdi,[]).append(dict(n=nmr,q=round(sum(q)/len(q),1),h=round(sum(h)/len(h),1) if h else 0,tipe=typ,lebar=leb,kond=kon))
OVR={'37 Karang Lebak':'Kr. Lebak','39 Sbr Titing':'Sumbertiting','42 Sbr Karang Labak':'Sbr. Karanglebak'}
# --- gabung
keys={norm(n):n for n in order}; out=[]; unm=[]
for k,(f,yr,nm) in sorted(best.items(),key=lambda x:x[1][2]):
    k={'ingas1dan2':'ingas','kedungpucang':'pucang'}.get(k,k)
    mk=difflib.get_close_matches(k,list(keys),1,0.8); xn=keys[mk[0]] if mk else None
    B,S,A=kml(f)
    for pre,x in OVR.items():
        if os.path.basename(f).startswith(pre): xn=x
    if not xn: unm.append(nm)
    dk=difflib.get_close_matches(norm(xn or nm),[norm(x) for x in deb],1,0.8)
    dd=[v for x,v in deb.items() if dk and norm(x)==dk[0]]; dkk=dk[0] if dk else ''
    out.append(dict(id=len(out)+1,src=os.path.basename(f),nama=(xn or nm),juru=dij.get(xn,'Belum terdaftar'),tahun=yr,
        tanam=tanam(xn) if xn else None,xk=xkeys(xn) if xn else [],ters=[{'n':c['_o'],**{k:round(c[k],2) for k in('baku','padi','tebu','palawija','bero','panen')}} for kk in (xkeys(xn) if xn else []) for c in (T5a.get(kk) or {'ch':[]})['ch'] if c['_n'].startswith('ters')][:60],dk=dkk,debit=list({x['n']:x for x in dd[0]}.values()) if dd else [],bakuKmz=round(sum(a[1] for a in A),1),
        panjangKm=round(sum(s[1] for s in S)/1000,2),bangunan=B,saluran=S,areal=A))
per=''
for r in r5[:10]:
    for c in r:
        if isinstance(c,str) and 'PERIODE' in c.upper(): per=re.sub(r'(?i)^\s*PERIODE\s*(Tanggal)?\s*','',c).strip()
os.makedirs(OUT,exist_ok=True); fo=os.path.join(OUT,'data.json')
json.dump(dict(periode=per or '-',juru=sorted(set(dij.values())),di=out),open(fo,'w'),separators=(',',':'))
print('data.json',round(os.path.getsize(fo)/1e6,2),'MB |',len(out),'DI | periode:',per)
print('DI tanpa data tanaman:',[d['nama'] for d in out if not d['tanam']])
print('DI tanpa debit:',[d['nama'] for d in out if not d['debit']])
print('Selisih baku Excel vs peta >15%:',[(d['nama'],d['tanam']['baku'],d['bakuKmz']) for d in out if d['tanam'] and d['tanam']['baku'] and abs(d['tanam']['baku']-d['bakuKmz'])/d['tanam']['baku']>.15])
